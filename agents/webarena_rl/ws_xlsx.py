"""mem_timeline.py --csv output -> an Excel workbook: sheet "data" (the CSV as numbers), sheet "figure" (a native
scatter-with-lines chart of every system, the figure's colours and axis titles), sheet "notes" (--note lines).
  <python with openpyxl> -I ws_xlsx.py --csv <csv> --out <xlsx> [--note <text> ...]"""
import argparse, csv

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font

ap = argparse.ArgumentParser()
ap.add_argument("--csv", required=True)
ap.add_argument("--out", required=True)
ap.add_argument("--note", action="append", default=[])
a = ap.parse_args()

COLOR = {"tfork (ours)": "2A78D6", "CRIU": "EB6834", "AgentENV": "1BAF7A", "no fork (fresh env per rollout)": "EDA100"}
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)

rows = list(csv.reader(open(a.csv)))
head, body = rows[0], rows[1:]
wb = Workbook()
ws = wb.active
ws.title = "data"
ws.append(head)
for r in body:
    ws.append([float(v) if v else None for v in r])
for c in ws[1]:
    c.font, c.alignment = BOLD, Alignment(wrap_text=True, vertical="top")
for col in ws.iter_cols(min_row=2):
    for c in col:
        c.font, c.number_format = FONT, "0.000" if c.column % 2 else "0.00"
for j in range(len(head)):
    ws.column_dimensions[ws.cell(1, j + 1).column_letter].width = 16
ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

ch = ScatterChart()
ch.style, ch.width, ch.height, ch.legend.position = 2, 22, 13, "t"
ch.x_axis.title, ch.y_axis.title = "time since the epoch started (min)", "environments' working set (GiB)"
ch.x_axis.scaling.min = ch.y_axis.scaling.min = 0
ch.x_axis.delete = ch.y_axis.delete = False   # openpyxl 3.1 otherwise hides both axes in Excel
ch.x_axis.majorGridlines = None
xmax = 0
for j in range(0, len(head), 2):
    name = head[j].rsplit(" time (min)", 1)[0]
    n = sum(1 for r in body if r[j])
    xmax = max(xmax, max(float(r[j]) for r in body if r[j]))
    s = Series(Reference(ws, min_col=j + 2, min_row=2, max_row=n + 1),
               Reference(ws, min_col=j + 1, min_row=2, max_row=n + 1), title=name)
    s.marker.symbol, s.smooth = "none", False
    s.graphicalProperties.line.solidFill = COLOR.get(name, "898781")
    s.graphicalProperties.line.width = 25400   # 2 pt
    ch.series.append(s)
ch.x_axis.scaling.max = round(xmax * 1.02, 1)
fig = wb.create_sheet("figure")
fig.add_chart(ch, "B2")

notes = wb.create_sheet("notes")
for line in a.note:
    notes.append([line])
for c in notes["A"]:
    c.font = FONT
notes.column_dimensions["A"].width = 120
wb.save(a.out)
print("wrote", a.out, f"({len(body)} rows, {len(ch.series)} series)")
